"""
Pruebas del rango de fechas del historial de indicadores.

El panel dejó de poder preguntar solo «los últimos treinta días» y ahora puede
preguntar «el 14 de agosto» o «del 1 al 30 de junio». Lo que se comprueba aquí
no es que el filtro filtre, sino que el rango signifique lo que la pantalla
dice: que un día pedido sea ese día y no el de al lado —que es el error que
aparece en cuanto alguien mezcla la hora de Trujillo con la de Greenwich—, que
las fechas se redondeen al período que las contiene en vez de partirlo, y que
el Excel salga con el mismo tramo que se está viendo en pantalla.
"""

from datetime import datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models.fraud_log import FraudLog
from app.models.order import Order, OrderStatus
from app.models.user import UserRole
from app.services import fraud_metrics_service
from app.services.errors import OperacionNoPermitida
from app.services.fraud_metrics_service import ZONA_DE_LA_TIENDA

from tests.conftest import cabeceras_de, crear_usuario


async def _evaluacion(sesion, usuario, cuando: datetime, monto: float = 100.0):
    """Una compra evaluada, fechada a mano."""
    orden = Order(
        user_id=usuario.id,
        total_amount=monto,
        status=OrderStatus.PENDING,
        shipping_address="Jr. Alfonso Ugarte 493",
        shipping_city="Trujillo",
        created_at=cuando,
    )
    sesion.add(orden)
    await sesion.flush()
    sesion.add(
        FraudLog(
            order_id=orden.id,
            fraud_score=0.2,
            decision="APPROVED",
            evaluated_at=cuando,
            detection_time_ms=2.0,
        )
    )
    await sesion.commit()
    return orden


# ── El tramo que se pide es el que se devuelve ───────────────────────────────


@pytest.mark.asyncio
async def test_un_solo_dia_devuelve_solo_ese_dia(sesion, ahora):
    usuario = await crear_usuario(sesion)
    ayer = ahora - timedelta(days=1)

    await _evaluacion(sesion, usuario, ayer, monto=500.0)
    await _evaluacion(sesion, usuario, ahora, monto=900.0)

    serie = await fraud_metrics_service.historial(
        sesion, "day", desde=ayer.date(), hasta=ayer.date()
    )

    assert len(serie) == 1
    assert serie[0].inicio == ayer.date()
    # La compra de hoy queda fuera: si se colara, el monto sería 1400.
    assert serie[0].evaluaciones == 1
    assert serie[0].monto_aprobado == 500.0


@pytest.mark.asyncio
async def test_el_rango_incluye_los_dos_extremos(sesion, ahora):
    usuario = await crear_usuario(sesion)
    for atras in (4, 2, 0):
        await _evaluacion(sesion, usuario, ahora - timedelta(days=atras))

    serie = await fraud_metrics_service.historial(
        sesion,
        "day",
        desde=(ahora - timedelta(days=4)).date(),
        hasta=ahora.date(),
    )

    # Cinco días, con los dos que limitan el rango dentro.
    assert len(serie) == 5
    assert serie[0].inicio == (ahora - timedelta(days=4)).date()
    assert serie[-1].inicio == ahora.date()
    assert sum(p.evaluaciones for p in serie) == 3


@pytest.mark.asyncio
async def test_una_compra_de_la_noche_entra_en_el_dia_de_trujillo(sesion, ahora):
    """
    A las diez de la noche en Trujillo ya son las tres de la mañana del día
    siguiente en UTC. Si el rango se resolviera en UTC, pedir «hoy» dejaría
    fuera las compras de la noche de hoy y las contaría en las de mañana.
    """
    usuario = await crear_usuario(sesion)
    esta_noche = ahora.replace(hour=22, minute=30)

    await _evaluacion(sesion, usuario, esta_noche)

    serie = await fraud_metrics_service.historial(
        sesion, "day", desde=ahora.date(), hasta=ahora.date()
    )

    assert len(serie) == 1
    assert serie[0].evaluaciones == 1


@pytest.mark.asyncio
async def test_las_fechas_se_redondean_al_periodo_que_las_contiene(sesion, ahora):
    """
    Pedir del 14 al 20 en escala mensual devuelve el mes entero, no un trozo.

    Devolver siete días rotulados «agosto» daría una tasa de detección que no
    es la de agosto, y es la que acabaría citada como tal.
    """
    usuario = await crear_usuario(sesion)
    primero_del_mes = ahora.replace(day=1, hour=9)

    await _evaluacion(sesion, usuario, primero_del_mes)

    dentro = ahora.replace(day=14).date()
    serie = await fraud_metrics_service.historial(
        sesion, "month", desde=dentro, hasta=ahora.replace(day=20).date()
    )

    assert len(serie) == 1
    assert serie[0].inicio == primero_del_mes.date()
    # La compra del día 1 cuenta, aunque el rango pedido empezaba el 14.
    assert serie[0].evaluaciones == 1


@pytest.mark.asyncio
async def test_sin_fecha_final_el_rango_llega_hasta_hoy(sesion, ahora):
    usuario = await crear_usuario(sesion)
    await _evaluacion(sesion, usuario, ahora)

    serie = await fraud_metrics_service.historial(
        sesion, "day", desde=(ahora - timedelta(days=2)).date()
    )

    assert len(serie) == 3
    assert serie[-1].inicio == ahora.date()
    assert sum(p.evaluaciones for p in serie) == 1


@pytest.mark.asyncio
async def test_una_fecha_final_futura_se_recorta_a_hoy(sesion, ahora):
    """
    No hay compras por venir: rellenar el gráfico con meses vacíos por delante
    haría parecer que la tienda dejó de vender.
    """
    usuario = await crear_usuario(sesion)
    await _evaluacion(sesion, usuario, ahora)

    serie = await fraud_metrics_service.historial(
        sesion,
        "day",
        desde=ahora.date(),
        hasta=(ahora + timedelta(days=30)).date(),
    )

    assert len(serie) == 1
    assert serie[-1].inicio == ahora.date()


# ── Lo que no se acepta ──────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_el_rango_al_reves_se_rechaza(sesion, ahora):
    await crear_usuario(sesion)

    with pytest.raises(OperacionNoPermitida):
        await fraud_metrics_service.historial(
            sesion,
            "day",
            desde=ahora.date(),
            hasta=(ahora - timedelta(days=5)).date(),
        )


@pytest.mark.asyncio
async def test_un_rango_enorme_se_rechaza_en_vez_de_recortarse(sesion, ahora):
    """
    Recortar en silencio daría un archivo que dice cubrir cinco años y trae los
    últimos 366 días. Es peor que un error, porque nadie lo nota.
    """
    await crear_usuario(sesion)

    with pytest.raises(OperacionNoPermitida):
        await fraud_metrics_service.historial(
            sesion,
            "day",
            desde=(ahora - timedelta(days=800)).date(),
            hasta=ahora.date(),
        )


@pytest.mark.asyncio
async def test_ese_mismo_rango_cabe_en_una_escala_mas_amplia(sesion, ahora):
    """La salida que ofrece el mensaje de error funciona de verdad."""
    await crear_usuario(sesion)

    serie = await fraud_metrics_service.historial(
        sesion,
        "month",
        desde=(ahora - timedelta(days=800)).date(),
        hasta=ahora.date(),
    )

    assert 26 <= len(serie) <= 28  # 800 días son unos 27 meses


# ── El rango viaja hasta la API y hasta el Excel ─────────────────────────────


@pytest.mark.asyncio
async def test_la_api_devuelve_el_rango_que_de_verdad_cubre(cliente, sesion, ahora):
    await crear_usuario(sesion, email="admin@ejemplo.com", rol=UserRole.ADMIN)
    cabeceras = await cabeceras_de(cliente, "admin@ejemplo.com")

    ayer = (ahora - timedelta(days=1)).date()
    respuesta = await cliente.get(
        f"/api/v1/fraud/history?granularity=day&start_date={ayer}&end_date={ayer}",
        headers=cabeceras,
    )

    assert respuesta.status_code == 200
    datos = respuesta.json()
    assert datos["range_start"] == str(ayer)
    assert datos["range_end"] == str(ayer)
    assert len(datos["periods"]) == 1


@pytest.mark.asyncio
async def test_en_escala_mensual_el_rango_llega_hasta_fin_de_mes(cliente, sesion, ahora):
    """
    `range_end` es el último día del período, no su primer día. Sin eso el
    reporte de un mes diría «del 01/09 al 01/09» y parecería de un solo día.
    """
    await crear_usuario(sesion, email="admin@ejemplo.com", rol=UserRole.ADMIN)
    cabeceras = await cabeceras_de(cliente, "admin@ejemplo.com")

    dia = ahora.date()
    respuesta = await cliente.get(
        f"/api/v1/fraud/history?granularity=month&start_date={dia}&end_date={dia}",
        headers=cabeceras,
    )

    datos = respuesta.json()
    assert datos["range_start"] == str(dia.replace(day=1))
    assert datos["range_end"].startswith(f"{dia.year:04d}-{dia.month:02d}-")
    assert datos["range_end"] >= str(dia)


@pytest.mark.asyncio
async def test_el_rango_al_reves_responde_400_y_no_revienta(cliente, sesion, ahora):
    await crear_usuario(sesion, email="admin@ejemplo.com", rol=UserRole.ADMIN)
    cabeceras = await cabeceras_de(cliente, "admin@ejemplo.com")

    hoy = ahora.date()
    antes = (ahora - timedelta(days=3)).date()
    respuesta = await cliente.get(
        f"/api/v1/fraud/history?granularity=day&start_date={hoy}&end_date={antes}",
        headers=cabeceras,
    )

    assert respuesta.status_code == 400
    assert "posterior" in respuesta.json()["detail"].lower()


@pytest.mark.asyncio
async def test_el_excel_trae_el_mismo_rango_que_el_panel(cliente, sesion, ahora):
    """
    El archivo y la pantalla tienen que hablar del mismo tramo. Si el rango se
    quedara en la pantalla, alguien exportaría creyendo que baja el día que
    está mirando y se llevaría los últimos doce meses.
    """
    await crear_usuario(sesion, email="admin@ejemplo.com", rol=UserRole.ADMIN)
    usuario = await crear_usuario(sesion)
    cabeceras = await cabeceras_de(cliente, "admin@ejemplo.com")

    ayer = ahora - timedelta(days=1)
    await _evaluacion(sesion, usuario, ayer, monto=700.0)
    await _evaluacion(sesion, usuario, ahora, monto=300.0)

    consulta = f"granularity=day&start_date={ayer.date()}&end_date={ayer.date()}"

    panel = (
        await cliente.get(f"/api/v1/fraud/history?{consulta}", headers=cabeceras)
    ).json()
    respuesta = await cliente.get(f"/api/v1/fraud/report.xlsx?{consulta}", headers=cabeceras)

    assert respuesta.status_code == 200
    # El rango va en el nombre del archivo: quien acumula reportes en una
    # carpeta los distingue por lo que miden, no por cuándo los descargó.
    assert str(ayer.date()) in respuesta.headers["content-disposition"]

    libro = load_workbook(BytesIO(respuesta.content))
    portada = libro["Resumen"]
    etiquetas = {
        portada.cell(row=f, column=1).value: portada.cell(row=f, column=2).value
        for f in range(1, 12)
    }
    assert etiquetas["Rango medido"] == ayer.strftime("%d/%m/%Y")

    serie = libro["Indicadores por período"]
    assert serie.max_row == 2  # el encabezado y un solo día
    assert serie.cell(row=2, column=2).value == panel["total_evaluations"] == 1
